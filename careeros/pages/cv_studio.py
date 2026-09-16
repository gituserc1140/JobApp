from io import BytesIO
import re

import streamlit as st
from docx import Document
from fpdf import FPDF
from pypdf import PdfReader

from careeros.services.openrouter import generate_text

st.title("🧾 CV Studio")


def extract_text(uploaded_file) -> str:
    name = uploaded_file.name.lower()
    if name.endswith(".txt"):
        return uploaded_file.getvalue().decode("utf-8", errors="ignore")
    if name.endswith(".pdf"):
        pdf = PdfReader(BytesIO(uploaded_file.getvalue()))
        return "\n".join(page.extract_text() or "" for page in pdf.pages)
    if name.endswith(".docx"):
        doc = Document(BytesIO(uploaded_file.getvalue()))
        return "\n".join(p.text for p in doc.paragraphs)
    return uploaded_file.getvalue().decode("utf-8", errors="ignore")


def to_pdf_bytes(text: str) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_font("Helvetica", size=10)
    for line in text.splitlines() or [text]:
        safe_line = re.sub(r"[^\x00-\x7F]+", " ", line)
        pdf.multi_cell(0, 6, safe_line)
    raw = pdf.output(dest="S")
    if isinstance(raw, (bytes, bytearray)):
        return bytes(raw)
    return raw.encode("latin-1")


def to_docx_bytes(text: str) -> bytes:
    doc = Document()
    for line in text.splitlines() or [text]:
        doc.add_paragraph(line)
    bio = BytesIO()
    doc.save(bio)
    return bio.getvalue()


cv_file = st.file_uploader("Upload CV", type=["txt", "pdf", "docx"])
if cv_file:
    cv_text = extract_text(cv_file)
    st.session_state["cv_text"] = cv_text
    st.success("CV parsed successfully.")
    st.text_area("Parsed CV", cv_text[:5000], height=200)

base_cv = st.session_state.get("cv_text", "")
role_type = st.selectbox("Generate CV version", ["Consulting", "Product Management", "UX", "Engineering"])

if st.button("Generate role-specific CV", type="primary"):
    if not base_cv.strip():
        st.warning("Upload a CV first.")
    else:
        prompt = f"Rewrite this CV for a {role_type} role. Keep it concise, outcome-focused, ATS-friendly:\n\n{base_cv[:12000]}"
        ai_text = generate_text(prompt)
        generated = ai_text or f"{role_type} CV Version\n\n{base_cv}"
        st.session_state["generated_cv"] = generated
        st.session_state["generated_cv_text"] = generated

if st.session_state.get("generated_cv"):
    if "generated_cv_text" not in st.session_state:
        st.session_state["generated_cv_text"] = st.session_state["generated_cv"]
    st.text_area("Generated CV", key="generated_cv_text", height=320)
    output = st.session_state["generated_cv_text"]

    st.download_button("Export PDF", data=to_pdf_bytes(output), file_name=f"cv_{role_type.lower().replace(' ', '_')}.pdf", mime="application/pdf")
    st.download_button("Export DOCX", data=to_docx_bytes(output), file_name=f"cv_{role_type.lower().replace(' ', '_')}.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
